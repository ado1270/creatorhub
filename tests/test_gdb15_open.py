"""gdb15 开放 API 客户端（CreatorHub 侧）回归测试。

盯的是四件事：
  1. 拉取用游标续传，既不重也不漏；
  2. 请求头走的是 x-open-api-key，不是推送用的那把 x-api-key；
  3. 服务端报错要翻译成客户看得懂的中话；
  4. 导出的表格不能被评论正文带跑（逗号串位、公式注入）。
"""
import asyncio
import csv
import io
import json
import tempfile
import unittest
from pathlib import Path

import httpx
from openpyxl import load_workbook

from app import db, gdb15_open
from app.settings import set_setting


def _lead(index: int, **kwargs) -> dict:
    lead = {
        "id": f"lead-{index}",
        "platform": "douyin",
        "nickname": "小明",
        "commentText": f"第 {index} 条评论",
        "sourceUrl": f"https://www.douyin.com/video/{index}",
        "sourceKeyword": "二手电动车",
        "intentKeyword": "多少钱",
        "intentLevel": "HIGH",
        "status": "PENDING",
        "capturedAt": "2026-10-29T02:00:00.000Z",
        "updatedAt": "2026-10-29T03:00:00.000Z",
        "owner": {"id": "u1", "nickname": "张经理", "email": "z@example.com"},
    }
    lead.update(kwargs)
    return lead


class OpenApiTestBase(unittest.TestCase):
    def setUp(self):
        self.previous_engine = db._engine
        self.tmp = tempfile.TemporaryDirectory()
        db.init_db(str(Path(self.tmp.name) / "openapi.db"))
        set_setting("license_url", "https://www.gdb15.com/api/v1/license")
        set_setting("license_api_key", "lk_test")

    def tearDown(self):
        if db._engine is not None:
            db._engine.dispose()
        db._engine = self.previous_engine
        self.tmp.cleanup()

    def _client(self, handler):
        return httpx.AsyncClient(transport=httpx.MockTransport(handler))

    def _collect(self, handler, **kwargs):
        async def run():
            async with self._client(handler) as client:
                return await gdb15_open.collect_leads(client=client, **kwargs)

        return asyncio.run(run())


class ConfigTest(OpenApiTestBase):
    def test_base_url_normalized(self):
        self.assertEqual(
            gdb15_open.normalize_base_url("https://www.gdb15.com"),
            "https://www.gdb15.com/api/v1/open")
        self.assertEqual(
            gdb15_open.normalize_base_url("https://www.gdb15.com/api/v1/open/"),
            "https://www.gdb15.com/api/v1/open")
        self.assertEqual(
            gdb15_open.normalize_base_url("www.gdb15.com"),
            "https://www.gdb15.com/api/v1/open")

    def test_base_url_rejects_insecure(self):
        with self.assertRaises(ValueError):
            gdb15_open.normalize_base_url("http://www.gdb15.com")
        with self.assertRaises(ValueError):
            gdb15_open.normalize_base_url("https://www.gdb15.com?x=1")

    def test_status_never_leaks_key(self):
        gdb15_open.save_open_api_settings(None, "gdb15_secret_key")
        status = gdb15_open.open_api_status()
        self.assertTrue(status["api_key_set"])
        self.assertTrue(status["configured"])
        self.assertNotIn("gdb15_secret_key", str(status))

    def test_empty_key_keeps_saved_one(self):
        gdb15_open.save_open_api_settings(None, "gdb15_secret_key")
        gdb15_open.save_open_api_settings(None, "")
        self.assertEqual(gdb15_open.open_api_key(), "gdb15_secret_key")

    def test_card_code_is_rejected_at_save_time(self):
        """卡密填进开放 API Key 框，要当场说清楚，不能等测试时回一句 401。"""
        with self.assertRaises(ValueError) as ctx:
            gdb15_open.save_open_api_settings(None, "JUFW-HDW6-TCK4-CDD4")
        self.assertIn("卡密", str(ctx.exception))

    def test_push_key_is_rejected_at_save_time(self):
        with self.assertRaises(ValueError) as ctx:
            gdb15_open.save_open_api_settings(
                None, "lk_c2a1cddde580d0dc9a4573934c9a1a327703b44ad151399e")
        self.assertIn("线索推送", str(ctx.exception))

    def test_rejected_key_does_not_half_save_the_url(self):
        """Key 填错就不该顺带把地址改了 —— 要么一起存，要么都不存。"""
        gdb15_open.save_open_api_settings("https://www.gdb15.com", "gdb15_ok")
        with self.assertRaises(ValueError):
            gdb15_open.save_open_api_settings("https://other.example.com",
                                              "JUFW-HDW6-TCK4-CDD4")
        self.assertEqual(gdb15_open.open_api_base(),
                         "https://www.gdb15.com/api/v1/open")

    def test_real_open_api_key_passes_shape_check(self):
        gdb15_open.save_open_api_settings(None, "gdb15_abc123")
        self.assertEqual(gdb15_open.open_api_key(), "gdb15_abc123")

    def test_site_base_follows_push_address(self):
        set_setting("license_url", "")
        set_setting("lead_push_url", "https://www.gdb15.com/api/v1/leads/webhook")
        self.assertEqual(gdb15_open.open_api_base(),
                         "https://www.gdb15.com/api/v1/open")


class CollectTest(OpenApiTestBase):
    def test_cursor_paging_keeps_every_lead_once(self):
        seen_cursors = []

        def handler(request: httpx.Request) -> httpx.Response:
            cursor = request.url.params.get("cursor")
            seen_cursors.append(cursor)
            if not cursor:
                return httpx.Response(200, json={
                    "items": [_lead(1), _lead(2)],
                    "nextCursor": "c1", "hasMore": True})
            if cursor == "c1":
                return httpx.Response(200, json={
                    "items": [_lead(3)],
                    "nextCursor": "c2", "hasMore": True})
            return httpx.Response(200, json={
                "items": [_lead(4)],
                "nextCursor": "c3", "hasMore": False})

        result = self._collect(handler,
                               base_url="https://www.gdb15.com/api/v1/open",
                               api_key="k")
        ids = [lead["id"] for lead in result["leads"]]
        self.assertEqual(ids, ["lead-1", "lead-2", "lead-3", "lead-4"])
        self.assertFalse(result["truncated"])
        # 第一页不带游标，之后每页都带服务端给的游标。
        self.assertEqual(seen_cursors, [None, "c1", "c2"])

    def test_uses_open_api_key_header_not_push_key(self):
        seen = {}

        def handler(request: httpx.Request) -> httpx.Response:
            seen["open"] = request.headers.get("x-open-api-key")
            seen["push"] = request.headers.get("x-api-key")
            return httpx.Response(200, json={"items": [], "hasMore": False})

        self._collect(handler, base_url="https://www.gdb15.com/api/v1/open",
                      api_key="open_key")
        self.assertEqual(seen["open"], "open_key")
        self.assertIsNone(seen["push"])

    def test_same_timestamp_leads_are_not_skipped(self):
        """同一毫秒进来的一批线索，用游标也要一条不落。"""
        def handler(request: httpx.Request) -> httpx.Response:
            cursor = request.url.params.get("cursor")
            if not cursor:
                return httpx.Response(200, json={
                    "items": [_lead(1), _lead(2)],
                    "nextCursor": "same", "hasMore": True})
            if cursor == "same":
                return httpx.Response(200, json={
                    "items": [_lead(3)],
                    "nextCursor": "same2", "hasMore": True})
            return httpx.Response(200, json={
                "items": [_lead(4)], "nextCursor": "same3", "hasMore": False})

        result = self._collect(handler,
                               base_url="https://www.gdb15.com/api/v1/open",
                               api_key="k")
        self.assertEqual(len(result["leads"]), 4)

    def test_repeated_cursor_stops_instead_of_looping(self):
        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(200, json={
                "items": [_lead(1)], "nextCursor": "same", "hasMore": True})

        result = self._collect(handler,
                               base_url="https://www.gdb15.com/api/v1/open",
                               api_key="k")
        # 服务端说还有下一页却把游标原地不动 —— 停在这里，别把同一页拉两遍。
        self.assertEqual(len(result["leads"]), 2)

    def test_max_leads_truncates(self):
        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(200, json={
                "items": [_lead(i) for i in range(10)],
                "nextCursor": "next", "hasMore": True})

        result = self._collect(handler,
                               base_url="https://www.gdb15.com/api/v1/open",
                               api_key="k", max_leads=3)
        self.assertEqual(len(result["leads"]), 3)
        self.assertTrue(result["truncated"])

    def test_filters_are_forwarded(self):
        seen = {}

        def handler(request: httpx.Request) -> httpx.Response:
            seen.update({k: v for k, v in request.url.params.items()})
            return httpx.Response(200, json={"items": [], "hasMore": False})

        self._collect(handler, base_url="https://www.gdb15.com/api/v1/open",
                      api_key="k", platform="douyin", status="PENDING",
                      source_keyword="电动车",
                      since="2026-10-01T00:00:00.000Z")
        self.assertEqual(seen["platform"], "douyin")
        self.assertEqual(seen["status"], "PENDING")
        self.assertEqual(seen["sourceKeyword"], "电动车")
        self.assertEqual(seen["since"], "2026-10-01T00:00:00.000Z")

    def test_401_is_translated(self):
        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(401, json={"message": "invalid api key"})

        with self.assertRaises(gdb15_open.OpenApiError) as ctx:
            self._collect(handler,
                          base_url="https://www.gdb15.com/api/v1/open", api_key="k")
        self.assertIn("吊销", str(ctx.exception))

    def test_401_with_card_code_names_the_card_code(self):
        """已经存进去的卡密，点测试时也要说人话，别只说「Key 无效」。"""
        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(401, json={"message": "API Key 无效"})

        with self.assertRaises(gdb15_open.OpenApiError) as ctx:
            self._collect(handler,
                          base_url="https://www.gdb15.com/api/v1/open",
                          api_key="JUFW-HDW6-TCK4-CDD4")
        self.assertIn("卡密", str(ctx.exception))

    def test_401_with_push_key_names_the_push_key(self):
        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(401, json={"message": "API Key 无效"})

        with self.assertRaises(gdb15_open.OpenApiError) as ctx:
            self._collect(handler,
                          base_url="https://www.gdb15.com/api/v1/open",
                          api_key="lk_c2a1cddde580d0dc9a4573934c9a1a327703b44ad151399e")
        self.assertIn("线索推送", str(ctx.exception))

    def test_403_is_translated(self):
        def handler(request: httpx.Request) -> httpx.Response:
            return httpx.Response(403, json={"message": "套餐已到期"})

        with self.assertRaises(gdb15_open.OpenApiError) as ctx:
            self._collect(handler,
                          base_url="https://www.gdb15.com/api/v1/open", api_key="k")
        self.assertIn("套餐已到期", str(ctx.exception))

    def test_missing_key_is_reported_before_any_request(self):
        def handler(request: httpx.Request) -> httpx.Response:
            raise AssertionError("不该发出请求")

        with self.assertRaises(gdb15_open.OpenApiError) as ctx:
            self._collect(handler,
                          base_url="https://www.gdb15.com/api/v1/open", api_key="")
        self.assertIn("Key", str(ctx.exception))


class ExportTest(OpenApiTestBase):
    def test_csv_keeps_columns_aligned(self):
        """评论正文里有逗号和引号，列不能串位。"""
        leads = [_lead(1, commentText='带"引号"和,逗号的正文')]
        payload = gdb15_open.build_csv(leads).decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(payload)))
        self.assertEqual(rows[0], gdb15_open.HEADERS)
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][2], '带"引号"和,逗号的正文')
        self.assertEqual(rows[1][0], "抖音")
        self.assertEqual(rows[1][6], "高")
        self.assertEqual(rows[1][7], "待跟进")

    def test_csv_has_bom_for_excel(self):
        payload = gdb15_open.build_csv([_lead(1)])
        self.assertTrue(payload.startswith(b"\xef\xbb\xbf"))

    def test_xlsx_escapes_formula(self):
        """评论正文以 = 开头时不能被 Excel 当成公式执行。"""
        payload = gdb15_open.build_xlsx([_lead(1, commentText="=1+1")])
        workbook = load_workbook(io.BytesIO(payload))
        sheet = workbook.active
        self.assertEqual(sheet.cell(row=2, column=3).value, "'=1+1")

    def test_xlsx_has_link_and_headers(self):
        payload = gdb15_open.build_xlsx([_lead(1)])
        workbook = load_workbook(io.BytesIO(payload))
        sheet = workbook.active
        self.assertEqual([c.value for c in sheet[1]], gdb15_open.HEADERS)
        self.assertEqual(sheet.cell(row=2, column=4).value,
                         "https://www.douyin.com/video/1")
        self.assertEqual(sheet.cell(row=2, column=4).hyperlink.target,
                         "https://www.douyin.com/video/1")

    def test_xlsx_note_row(self):
        payload = gdb15_open.build_xlsx([_lead(1)], note="这是说明")
        workbook = load_workbook(io.BytesIO(payload))
        sheet = workbook.active
        self.assertEqual(sheet.cell(row=1, column=1).value, "这是说明")

    def test_time_is_local_not_utc(self):
        """时间要写本地时间，否则客户会以为线索早了八小时。"""
        payload = gdb15_open.build_xlsx([_lead(1)])
        workbook = load_workbook(io.BytesIO(payload))
        sheet = workbook.active
        captured = sheet.cell(row=2, column=10).value
        self.assertNotIn("Z", str(captured))
        self.assertIn(":", str(captured))


class PushStatusTest(OpenApiTestBase):
    def test_status_writeback(self):
        seen = {}

        def handler(request: httpx.Request) -> httpx.Response:
            seen["method"] = request.method
            seen["path"] = request.url.path
            seen["key"] = request.headers.get("x-open-api-key")
            seen["body"] = json.loads(request.content.decode("utf-8"))
            return httpx.Response(200, json={"id": "lead-1", "status": "CONTACTED"})

        async def run():
            async with self._client(handler) as client:
                return await gdb15_open.push_status(
                    "lead-1", "CONTACTED", "已电话联系",
                    base_url="https://www.gdb15.com/api/v1/open", api_key="k",
                    client=client)

        result = asyncio.run(run())
        self.assertEqual(result["status"], "CONTACTED")
        self.assertEqual(seen["method"], "PUT")
        self.assertTrue(seen["path"].endswith("/leads/lead-1/status"))
        self.assertEqual(seen["key"], "k")
        self.assertEqual(seen["body"]["status"], "CONTACTED")
        self.assertEqual(seen["body"]["note"], "已电话联系")

    def test_invalid_status_is_rejected(self):
        async def run():
            return await gdb15_open.push_status(
                "lead-1", "NOPE",
                base_url="https://www.gdb15.com/api/v1/open", api_key="k")

        with self.assertRaises(gdb15_open.OpenApiError):
            asyncio.run(run())


if __name__ == "__main__":
    unittest.main()
